from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any, Iterable
import csv
import json

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CSV_FIELDS = ["index", "name", "identical", "name_matches", "pages_before", "pages_after", "page_delta", "size_delta_bytes", "tokens_added", "tokens_deleted", "text_change_percent", "technical_change_percent", "severity", "level", "ocr_pages_before", "ocr_pages_after", "ollama_status", "ollama_prompt_tokens", "ollama_completion_tokens", "ollama_total_tokens", "ollama_summary"]

NAVY = "17365D"
TEAL = "0F6B6D"
LIGHT_BLUE = "D9EAF7"
LIGHT_TEAL = "DDEBF0"
LIGHT_GRAY = "F2F4F6"
DARK_TEXT = "263238"
MUTED_TEXT = "5F6B73"
WHITE = "FFFFFF"
AMBER = "FCE4B2"
AMBER_TEXT = "7F6000"
RED = "F4CCCC"
RED_TEXT = "9C0006"
GREEN = "D9EAD3"
GREEN_TEXT = "274E13"
THIN_GRAY = Side(style="thin", color="D8DEE3")


def flatten(item: dict[str, Any]) -> dict[str, Any]:
    ai = item.get("ollama", {})
    summary_value = ai.get("summary", {}).get("summary", "")
    summary = summary_value.get("text", "") if isinstance(summary_value, dict) else summary_value
    usage = ai.get("usage", {})
    return {
        "index": item["index"], "name": item["name_before"], "identical": item["identical"], "name_matches": item["name_matches"],
        "pages_before": item["before"]["page_count"], "pages_after": item["after"]["page_count"], "page_delta": item["page_delta"],
        "size_delta_bytes": item["size_delta_bytes"], "tokens_added": item["text_diff"]["tokens_added"], "tokens_deleted": item["text_diff"]["tokens_deleted"],
        "text_change_percent": item["text_diff"]["change_percent"], "technical_change_percent": item["scores"]["technical_change_percent"],
        "severity": item["scores"]["severity"], "level": item["scores"]["level"],
        "ocr_pages_before": ",".join(map(str, item["before"]["pages_requiring_ocr"])), "ocr_pages_after": ",".join(map(str, item["after"]["pages_requiring_ocr"])),
        "ollama_status": ai.get("status", "not_requested"),
        "ollama_prompt_tokens": int(usage.get("prompt_tokens", 0) or 0),
        "ollama_completion_tokens": int(usage.get("completion_tokens", 0) or 0),
        "ollama_total_tokens": int(usage.get("total_tokens", 0) or 0),
        "ollama_summary": summary or "",
    }


def _title(sheet, title: str, subtitle: str, last_column: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(name="Arial", size=16, bold=True, color=NAVY)
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 27
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(2, 1, subtitle)
    sheet.cell(2, 1).font = Font(name="Arial", size=10, italic=True, color=MUTED_TEXT)
    sheet.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 27
    for cell in sheet[3][:last_column]:
        cell.fill = PatternFill("solid", fgColor=TEAL)
    sheet.row_dimensions[3].height = 4


def _header(row: Iterable[Any]) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color=WHITE))


def _body(sheet, min_row: int, max_row: int, max_col: int, wrap_columns: set[int] | None = None) -> None:
    wrap_columns = wrap_columns or set()
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=DARK_TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_columns)
            cell.border = Border(bottom=THIN_GRAY)
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F8FAFB")


def _set_widths(sheet, widths: list[float]) -> None:
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _priority_label(score: float) -> str:
    if score <= 10: return "Mínima"
    if score <= 25: return "Baja"
    if score <= 45: return "Moderada"
    if score <= 65: return "Alta"
    if score <= 85: return "Muy alta"
    return "Crítica"


def _ollama_text(item: dict[str, Any]) -> str:
    value = item.get("ollama", {}).get("summary", {}).get("summary", "")
    return value.get("text", "") if isinstance(value, dict) else str(value or "")


def _evidence_ids(value: Any) -> str:
    if not isinstance(value, list): return ""
    return ", ".join(str(item) for item in value)


def _findings(comparisons: list[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    def add(index: Any, name: str, finding_type: str, description: str, evidence: str, severity: float, source: str, page_before: Any = "", page_after: Any = "", amount_before: str = "", amount_after: str = "", note: str = "Revisar contra los documentos fuente; no se atribuye causa.") -> None:
        rows.append([index, name, finding_type, page_before, page_after, description, amount_before, amount_after, evidence, severity, _priority_label(severity), source, note])

    labels = {
        "amounts": "Apariciones de importes", "dates": "Apariciones de fechas",
        "identifiers": "Apariciones de identificadores", "diagnostic_codes": "Apariciones de códigos documentales",
        "signature_terms": "Referencias textuales de firma",
    }
    for item in comparisons:
        index, name = item["index"], item["name_before"]
        severity = float(item["scores"]["severity"])
        if item["page_delta"]:
            add(index, name, "Estructura documental", f"Diferencia neta de páginas: {item['page_delta']:+d}.", "E_PAGE_DELTA", severity, "Comparación determinística")
        text_change = float(item["text_diff"]["change_percent"])
        if text_change:
            add(index, name, "Texto extraído", f"El texto extraído presenta {text_change:.1f}% de diferencia. La métrica no determina por sí sola el sentido documental del cambio.", "E_TEXT_CHANGE", severity, "Comparación determinística")
        for money in item.get("monetary_line_changes", []):
            add(
                index, name, "Diferencia de importe en línea documental", money.get("description", ""), money["evidence_id"], severity,
                "Comparación monetaria por línea", money.get("page_before") or "", money.get("page_after") or "",
                ", ".join(money.get("amounts_before", [])), ", ".join(money.get("amounts_after", [])), money.get("review_note", "Requiere revisión contra los documentos fuente."),
            )
        for key, label in labels.items():
            values = item.get("sensitive_changes", {}).get(key, {})
            added, removed = int(values.get("added_count", 0)), int(values.get("removed_count", 0))
            if added or removed:
                description = f"{label}: {added} apariciones agregadas y {removed} eliminadas. El conteo no confirma correspondencia campo a campo."
                add(index, name, label, description, f"E_{key.upper()}", severity, "Detección por patrón")
        changed_metadata = len(item.get("metadata_changes", {}))
        if changed_metadata:
            add(index, name, "Metadatos PDF", f"Se detectaron diferencias en {changed_metadata} campos de metadatos.", "E_METADATA", severity, "Comparación determinística")
        ocr_applied = len(item["before"].get("pages_ocr_applied", [])) + len(item["after"].get("pages_ocr_applied", []))
        ocr_pending = len(item["before"].get("pages_requiring_ocr", [])) + len(item["after"].get("pages_requiring_ocr", []))
        if ocr_applied or ocr_pending:
            add(index, name, "OCR", f"OCR aplicado en {ocr_applied} páginas; {ocr_pending} páginas permanecen señaladas para revisión de legibilidad.", "E_OCR", severity, "Tesseract")
        ai = item.get("ollama", {})
        summary = ai.get("summary", {})
        for group, source in (("main_changes", "Resumen Ollama"), ("alerts", "Alerta Ollama")):
            for observation in summary.get(group, []) if isinstance(summary, dict) else []:
                if isinstance(observation, dict) and observation.get("text"):
                    add(index, name, source, observation["text"], _evidence_ids(observation.get("evidence_ids")), severity, f"Ollama: {ai.get('status', 'sin estado')}")
    if not rows:
        add("", "", "Sin hallazgos", "No se registraron diferencias con las reglas disponibles.", "", 0, "Comparación determinística", note="")
    return rows


def _summary_sheet(workbook: Workbook, month: str, manifest: dict[str, Any], comparisons: list[dict[str, Any]]) -> None:
    sheet = workbook.active
    sheet.title = "Resumen ejecutivo"
    generated = str(manifest.get("generated_at", ""))
    _title(sheet, "Resumen de comparación documental", f"Período {month} | Ejecución {manifest.get('run_id', 'sin identificador')} | Generado {generated}", 8)
    pages_before = sum(item["before"]["page_count"] for item in comparisons)
    pages_after = sum(item["after"]["page_count"] for item in comparisons)
    ocr_applied = sum(len(item[side].get("pages_ocr_applied", [])) for item in comparisons for side in ("before", "after"))
    ocr_pending = sum(len(item[side].get("pages_requiring_ocr", [])) for item in comparisons for side in ("before", "after"))
    max_score = max((float(item["scores"]["severity"]) for item in comparisons), default=0.0)
    ollama_counts = Counter(item.get("ollama", {}).get("status", "no solicitado") for item in comparisons)
    scope = f"{len(comparisons)} expediente seleccionado" if len(comparisons) == 1 else f"{len(comparisons)} expedientes incluidos"
    sheet["A5"], sheet["A6"] = "Período", month
    sheet["C5"], sheet["C6"] = "Alcance", scope
    sheet["E5"], sheet["E6"] = "Prioridad máxima", f"{max_score:.1f}/100 - {_priority_label(max_score)}"
    sheet["G5"], sheet["G6"] = "Estado del resumen", ", ".join(f"{key}: {value}" for key, value in sorted(ollama_counts.items())) or "No solicitado"
    for anchor in ("A5", "C5", "E5", "G5"):
        sheet[anchor].font = Font(name="Arial", size=9, bold=True, color=MUTED_TEXT)
    for anchor in ("A6", "C6", "E6", "G6"):
        sheet[anchor].font = Font(name="Arial", size=12, bold=True, color=NAVY)
        sheet[anchor].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("A6:B7"); sheet.merge_cells("C6:D7"); sheet.merge_cells("E6:F7"); sheet.merge_cells("G6:H7")
    metrics = [
        ("Páginas origen", pages_before), ("Páginas modificadas", pages_after),
        ("Diferencia neta", pages_after - pages_before), ("OCR aplicado", ocr_applied),
        ("OCR pendiente", ocr_pending), ("Expedientes con diferencia", sum(not item["identical"] for item in comparisons)),
    ]
    sheet["A9"] = "Indicadores de la ejecución"
    sheet["A9"].font = Font(name="Arial", size=11, bold=True, color=TEAL)
    for position, (label, value) in enumerate(metrics):
        column = 1 + (position % 3) * 2
        row = 10 + (position // 3) * 3
        sheet.cell(row, column, label)
        sheet.cell(row + 1, column, value)
        sheet.merge_cells(start_row=row, start_column=column, end_row=row, end_column=column + 1)
        sheet.merge_cells(start_row=row + 1, start_column=column, end_row=row + 1, end_column=column + 1)
        sheet.cell(row, column).fill = PatternFill("solid", fgColor=LIGHT_TEAL)
        sheet.cell(row, column).font = Font(name="Arial", size=9, bold=True, color=TEAL)
        sheet.cell(row + 1, column).font = Font(name="Arial", size=15, bold=True, color=NAVY)
        sheet.cell(row + 1, column).number_format = "#,##0;[Red]-#,##0"
    sheet["A16"] = "Lectura administrativa"
    sheet["A16"].font = Font(name="Arial", size=11, bold=True, color=TEAL)
    sheet.merge_cells("A17:H19")
    if len(comparisons) == 1:
        item = comparisons[0]
        text = _ollama_text(item) or f"El expediente {item['index']} registra una prioridad de revisión de {float(item['scores']['severity']):.1f}/100."
    else:
        text = "El informe ordena diferencias documentales para revisión administrativa. Los puntajes no determinan corrección clínica, causa ni responsabilidad."
    sheet["A17"] = text
    sheet["A17"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A17"].font = Font(name="Arial", size=10, color=DARK_TEXT)
    sheet["A21"] = "Leyenda de prioridad"
    sheet["A21"].font = Font(name="Arial", size=10, bold=True, color=NAVY)
    legend = [("Mínima/Baja", "0-25", GREEN), ("Moderada/Alta", ">25-65", AMBER), ("Muy alta/Crítica", ">65-100", RED)]
    for offset, (label, range_text, color) in enumerate(legend):
        column = 1 + offset * 2
        sheet.cell(22, column, label); sheet.cell(22, column + 1, range_text)
        sheet.cell(22, column).fill = PatternFill("solid", fgColor=color)
        sheet.cell(22, column).font = Font(name="Arial", size=9, bold=True, color=DARK_TEXT)
    sheet.merge_cells("A24:H25")
    sheet["A24"] = "Uso: priorización de revisión documental. No sustituye la lectura del expediente ni confirma diagnóstico, corrección, causa, manipulación, fraude, autoría o responsabilidad."
    sheet["A24"].font = Font(name="Arial", size=9, italic=True, color=MUTED_TEXT)
    sheet["A24"].alignment = Alignment(wrap_text=True, vertical="top")
    _set_widths(sheet, [17, 15, 17, 15, 17, 15, 19, 20])
    sheet.print_area = "A1:H25"
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1


def _comparison_sheet(workbook: Workbook, month: str, comparisons: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("Comparativo")
    _title(sheet, "Comparativo por expediente", f"Período {month} | Vista administrativa de diferencias observadas", 12)
    headers = ["Índice", "Expediente", "Páginas origen", "Páginas modificadas", "Diferencia", "Cambio texto", "Cambio técnico", "Prioridad", "Nivel", "OCR aplicado", "OCR pendiente", "Estado resumen"]
    sheet.append(headers)
    for item in comparisons:
        applied = len(item["before"].get("pages_ocr_applied", [])) + len(item["after"].get("pages_ocr_applied", []))
        pending = len(item["before"].get("pages_requiring_ocr", [])) + len(item["after"].get("pages_requiring_ocr", []))
        sheet.append([item["index"], item["name_before"], item["before"]["page_count"], item["after"]["page_count"], item["page_delta"], float(item["text_diff"]["change_percent"]) / 100, float(item["scores"]["technical_change_percent"]) / 100, float(item["scores"]["severity"]), _priority_label(float(item["scores"]["severity"])), applied, pending, item.get("ollama", {}).get("status", "No solicitado")])
    _header(sheet[4])
    _body(sheet, 5, sheet.max_row, len(headers), {2, 12})
    sheet.freeze_panes = "C5"; sheet.auto_filter.ref = f"A4:L{sheet.max_row}"
    _set_widths(sheet, [9, 34, 13, 16, 12, 14, 15, 12, 15, 13, 13, 18])
    for row in range(5, sheet.max_row + 1):
        sheet.cell(row, 6).number_format = "0.0%"; sheet.cell(row, 7).number_format = "0.0%"; sheet.cell(row, 8).number_format = "0.0"
    if sheet.max_row >= 5:
        sheet.conditional_formatting.add(f"H5:H{sheet.max_row}", ColorScaleRule(start_type="num", start_value=0, start_color="63BE7B", mid_type="num", mid_value=50, mid_color="FFEB84", end_type="num", end_value=100, end_color="F8696B"))
        sheet.conditional_formatting.add(f"K5:K{sheet.max_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=AMBER)))
    sheet.page_setup.orientation = "landscape"; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.page_setup.fitToWidth = 1


def _findings_sheet(workbook: Workbook, month: str, comparisons: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("Hallazgos")
    _title(sheet, "Hallazgos documentales", f"Período {month} | Observaciones respaldadas por identificadores de evidencia", 13)
    headers = ["Índice", "Expediente", "Tipo", "Página origen", "Página auditor", "Descripción", "Importes origen", "Importes auditor", "Evidence IDs", "Prioridad", "Nivel", "Fuente", "Nota de revisión"]
    sheet.append(headers)
    for row in _findings(comparisons): sheet.append(row)
    _header(sheet[4]); _body(sheet, 5, sheet.max_row, 13, {2, 3, 6, 7, 8, 9, 12, 13})
    sheet.freeze_panes = "A5"; sheet.auto_filter.ref = f"A4:M{sheet.max_row}"
    _set_widths(sheet, [9, 29, 28, 13, 13, 56, 34, 34, 25, 11, 14, 24, 48])
    for row in range(5, sheet.max_row + 1):
        sheet.cell(row, 10).number_format = "0.0"
        sheet.row_dimensions[row].height = 48
    if sheet.max_row >= 5:
        sheet.conditional_formatting.add(f"J5:J{sheet.max_row}", ColorScaleRule(start_type="num", start_value=0, start_color="63BE7B", mid_type="num", mid_value=50, mid_color="FFEB84", end_type="num", end_value=100, end_color="F8696B"))
    sheet.page_setup.orientation = "landscape"; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.page_setup.fitToWidth = 1


def _methodology_sheet(workbook: Workbook, month: str, manifest: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Metodología")
    _title(sheet, "Metodología y límites", f"Período {month} | Criterios de lectura del informe", 4)
    headers = ["Tema", "Definición", "Uso administrativo", "Limitación"]
    sheet.append(headers)
    rows = [
        ["Emparejamiento", "Prefijo numérico antes de ' - ' en ambos directorios.", "Relaciona las dos versiones del mismo índice.", "Los índices faltantes o duplicados requieren revisión de carga."],
        ["Cambio de texto", "Proporción de tokens agregados y eliminados sobre ambos textos extraídos.", "Prioriza diferencias de contenido extraíble.", "El orden de extracción y OCR pueden afectar el porcentaje."],
        ["Cambio técnico", "Combinación reproducible de páginas, texto, geometría, señales documentales y metadatos.", "Ordena expedientes para revisión.", "No mide corrección clínica ni validez documental."],
        ["Prioridad 0-100", "Indicador de triage basado en las métricas determinísticas.", "Define el orden sugerido de revisión.", "No confirma causa, intención, manipulación, fraude, autoría o responsabilidad."],
        ["OCR", "Tesseract se aplica a páginas con texto extraíble insuficiente.", "Reduce páginas sin contenido legible para el motor.", "Puede omitir o confundir caracteres; requiere revisión visual."],
        ["Señales por patrón", "Conteos de apariciones de importes, fechas, identificadores, códigos y referencias de firma.", "Señala categorías que merecen inspección.", "No confirma correspondencia campo a campo ni significado clínico."],
        ["Ollama", "Redacta desde métricas agregadas y evidence_ids calculados por Python.", "Facilita una lectura breve.", "No calcula diferencias, puntajes ni conclusiones clínicas."],
        ["P", "Cambio proporcional del número de páginas.", "Componente estructural.", "No describe el contenido de las páginas."],
        ["T", "Porcentaje de cambio del texto extraído.", "Componente textual.", "Depende de extracción/OCR."],
        ["V", "Cambio de geometría o rotación de páginas.", "Componente visual técnico disponible en servidor.", "No es una comparación píxel a píxel."],
        ["F", "Presencia de diferencias en señales documentales detectadas por patrón.", "Componente de priorización.", "Puede incluir falsos positivos."],
        ["M", "Presencia de diferencias en metadatos PDF.", "Componente técnico.", "Un cambio de metadatos no prueba cambio de contenido."],
    ]
    for row in rows: sheet.append(row)
    _header(sheet[4]); _body(sheet, 5, sheet.max_row, 4, {1, 2, 3, 4})
    sheet.freeze_panes = "A5"; sheet.auto_filter.ref = f"A4:D{sheet.max_row}"
    _set_widths(sheet, [23, 62, 43, 58])
    for row in range(5, sheet.max_row + 1): sheet.row_dimensions[row].height = 45
    start = sheet.max_row + 2
    sheet.cell(start, 1, "Limitaciones declaradas en la ejecución")
    sheet.cell(start, 1).font = Font(name="Arial", size=11, bold=True, color=TEAL)
    for offset, value in enumerate(manifest.get("limitations", []), 1):
        sheet.cell(start + offset, 1, value); sheet.merge_cells(start_row=start + offset, start_column=1, end_row=start + offset, end_column=4)
        sheet.cell(start + offset, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(start + offset, 1).font = Font(name="Arial", size=10, color=DARK_TEXT)
    sheet.page_setup.orientation = "landscape"; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.page_setup.fitToWidth = 1


def _execution_sheet(workbook: Workbook, month: str, manifest: dict[str, Any], comparisons: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("Ejecución")
    _title(sheet, "Trazabilidad de la ejecución", f"Período {month} | Parámetros, consumo y referencias técnicas", 7)
    usage = Counter()
    statuses = Counter()
    models = set()
    for item in comparisons:
        ai = item.get("ollama", {}); statuses[ai.get("status", "no solicitado")] += 1
        if ai.get("model"): models.add(str(ai["model"]))
        for key, value in ai.get("usage", {}).items(): usage[key] += int(value or 0)
    info = [
        ("Run ID", manifest.get("run_id", "")), ("Generado UTC", manifest.get("generated_at", "")),
        ("Período", month), ("Alcance", manifest.get("scope", "")),
        ("Índices solicitados", ", ".join(map(str, manifest.get("requested_indices") or [])) or "Mes completo"),
        ("Pares procesados", manifest.get("pair_count", len(comparisons))), ("Fallos", manifest.get("failure_count", 0)),
        ("Modelo Ollama", ", ".join(sorted(models)) or "No solicitado/disponible"),
        ("Estados Ollama", ", ".join(f"{key}: {value}" for key, value in sorted(statuses.items())) or "No solicitado"),
        ("Tokens prompt", usage["prompt_tokens"]), ("Tokens respuesta", usage["completion_tokens"]), ("Tokens totales", usage["total_tokens"]),
        ("Páginas OCR aplicadas", sum(len(item[side].get("pages_ocr_applied", [])) for item in comparisons for side in ("before", "after"))),
        ("Páginas OCR pendientes", sum(len(item[side].get("pages_requiring_ocr", [])) for item in comparisons for side in ("before", "after"))),
    ]
    sheet["A5"], sheet["B5"] = "Dato", "Valor"
    _header(sheet[5][:2])
    for row, (label, value) in enumerate(info, 6):
        sheet.cell(row, 1, label); sheet.cell(row, 2, value)
        sheet.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=NAVY)
        sheet.cell(row, 2).font = Font(name="Arial", size=10, color=DARK_TEXT)
        sheet.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    technical_start = 22
    sheet.cell(technical_start, 1, "Referencias técnicas por expediente")
    sheet.cell(technical_start, 1).font = Font(name="Arial", size=11, bold=True, color=TEAL)
    headers = ["Índice", "Archivo origen", "SHA-256 origen", "Archivo modificado", "SHA-256 modificado", "Bytes origen", "Bytes modificados"]
    for col, header in enumerate(headers, 1): sheet.cell(technical_start + 1, col, header)
    _header(sheet[technical_start + 1])
    for item in comparisons:
        sheet.append([item["index"], item["before"].get("filename", ""), item["before"].get("sha256", ""), item["after"].get("filename", ""), item["after"].get("sha256", ""), item["before"].get("size_bytes", 0), item["after"].get("size_bytes", 0)])
    _body(sheet, technical_start + 2, sheet.max_row, 7, {2, 3, 4, 5})
    sheet.freeze_panes = "A6"; sheet.auto_filter.ref = f"A{technical_start + 1}:G{sheet.max_row}"
    _set_widths(sheet, [23, 39, 68, 39, 68, 17, 18])
    for row in range(technical_start + 2, sheet.max_row + 1):
        sheet.cell(row, 6).number_format = "#,##0"; sheet.cell(row, 7).number_format = "#,##0"
    sheet.page_setup.orientation = "landscape"; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.page_setup.fitToWidth = 1


def write_reports(month: str, output_dir: Path, manifest: dict[str, Any]) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "comparison_details.json"
    csv_path = output_dir / "comparison_summary.csv"
    xlsx_path = output_dir / "comparison_summary.xlsx"
    json_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    comparisons = manifest.get("comparisons", [])
    rows = [flatten(item) for item in comparisons]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=CSV_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    workbook = Workbook()
    workbook.properties.title = f"Comparación documental {month}"
    workbook.properties.subject = "Priorización administrativa de diferencias documentales"
    workbook.properties.creator = "CompareForms"
    _summary_sheet(workbook, month, manifest, comparisons)
    _comparison_sheet(workbook, month, comparisons)
    _findings_sheet(workbook, month, comparisons)
    _methodology_sheet(workbook, month, manifest)
    _execution_sheet(workbook, month, manifest, comparisons)
    workbook.active = 0
    workbook.save(xlsx_path)

    expected = ["Resumen ejecutivo", "Comparativo", "Hallazgos", "Metodología", "Ejecución"]
    checked = load_workbook(xlsx_path, read_only=True, data_only=False)
    assert checked.sheetnames == expected
    assert checked["Resumen ejecutivo"]["A1"].value == "Resumen de comparación documental"
    assert checked["Comparativo"].max_column == 12
    checked.close()
    return {"json": str(json_path), "csv": str(csv_path), "xlsx": str(xlsx_path)}
