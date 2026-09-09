"""Auditor-facing XLSX snapshots, generated entirely in Python on the server.

openpyxl is deliberate: production does not have Node.js. Every long evidence
value is split into numbered continuation rows, never silently cut to Excel's
cell limit. Reviews in an exported workbook do not mutate portal records.
"""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
import math
import os
import re
from urllib.parse import urlparse
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

TYPE_LABELS = {"added": "Añadido", "removed": "Eliminado", "modified": "Modificado", "relocated": "Reubicado", "review": "Revisión"}
STATUS_LABELS = {"with_differences": "Con diferencias detectadas", "no_differences_detected": "Sin diferencias detectadas", "inconclusive": "No concluyente",
                 "completed": "Completado", "partial": "Parcial / no concluyente", "failed": "Fallido", "queued": "Pendiente", "running": "En proceso", "cancelled": "Cancelado"}
MAP_LABELS = {"matched": "Correspondencia", "identical": "Idéntica", **TYPE_LABELS}
COLORS = {"Añadido": "DDEBF7", "Eliminado": "FCE4D6", "Modificado": "FFF2CC", "Reubicado": "E4DFEC", "Revisión": "FFE3A3"}
NAVY, BLUE, GRAY = "193B5C", "DDEBF7", "526477"
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEADER_ROW = 6
MAX_FRAGMENT = 600


def _safe(value):
    if not isinstance(value, str):
        return value
    value = INVALID_XML.sub("�", value)
    if value.lstrip().startswith(("=", "+", "-", "@")):
        value = "'" + value
    if len(value) > 32767:
        raise ValueError("Texto demasiado largo para Excel; debe dividirse en continuaciones.")
    return value


def _parts(value: str | None) -> list[str]:
    value = str(value or "")
    # Preserve every character, including line breaks, in ordered fragments.
    return [value[i:i + MAX_FRAGMENT] for i in range(0, len(value), MAX_FRAGMENT)] or [""]


def _page(value) -> int | str:
    if value is None:
        return "Sin equivalente"
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValueError("Las páginas de evidencia deben ser enteros positivos o nulas.")
    return value


def _base(wb: Workbook, name: str, title: str, context: str, headers: list[str], widths: list[float]):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.sheet_properties.tabColor = NAVY if name in ("Resumen ejecutivo", "Hallazgos") else None
    ws.cell(2, 1, _safe(title)).font = Font(name="Arial", size=16, bold=True, color=NAVY)
    ws.cell(3, 1, _safe(context)).font = Font(name="Arial", size=10, italic=True, color=GRAY)
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 18
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(HEADER_ROW, col, header)
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color="FFFFFF"))
        ws.cell(4, col).border = Border(bottom=Side(style="thin", color="829EB7"))
    ws.row_dimensions[HEADER_ROW].height = 34
    ws.freeze_panes = "B7"
    ws.print_title_rows = "6:6"
    ws.print_options.horizontalCentered = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = "CompareForms - Página &P de &N"
    return ws


def _append(ws, row: list):
    target = max(HEADER_ROW + 1, ws.max_row + 1)
    needed_lines = 1
    for col, value in enumerate(row, 1):
        cell = ws.cell(target, col, _safe(value))
        cell.font = Font(name="Arial", size=10, color="1F2937")
        cell.alignment = Alignment(horizontal="left" if isinstance(value, str) else "right", vertical="top", wrap_text=True)
        if isinstance(value, str):
            width = ws.column_dimensions[get_column_letter(col)].width
            needed_lines = max(needed_lines, sum(max(1, math.ceil(len(part) / max(8, width - 3))) for part in value.split("\n")))
    ws.row_dimensions[target].height = max(30, needed_lines * 14 + 8)
    return target


def _finish(ws, table_name: str):
    count = ws.max_column
    last = ws.max_row
    if last >= HEADER_ROW + 1:
        tab = Table(displayName=table_name, ref=f"A{HEADER_ROW}:{get_column_letter(count)}{last}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)
    else:
        ws.cell(HEADER_ROW + 1, 1, "No hay registros para esta sección.").font = Font(name="Arial", size=10, italic=True, color=GRAY)
    ws.print_area = f"A1:{get_column_letter(count)}{ws.max_row}"


def _public_base() -> str | None:
    base = os.getenv("COMPAREFORMS_PUBLIC_URL", "").rstrip("/")
    parts = urlparse(base)
    # A URL is configuration, not a value extracted from the PDFs.
    return base if parts.scheme in {"https", "http"} and parts.netloc and not parts.username and not parts.password and not parts.query and not parts.fragment else None


def _link(ws, row: int, col: int, doc_id, page: int | None):
    base = _public_base()
    if not base or not doc_id or page is None:
        return
    try:
        identifier = str(UUID(str(doc_id)))
    except ValueError:
        return
    cell = ws.cell(row, col)
    cell.hyperlink = f"{base}/api/v1/documents/{identifier}/content#page={page}"
    cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")


def write_report(path: Path, *, period: str, run_id: str, cases: list[dict], status: str) -> Path:
    """Write a versioned snapshot for administrators; reject an empty comparison."""
    if not cases:
        raise ValueError("NO_COMPARABLE_PAIRS: no se genera un informe normal con cero expedientes.")
    if not re.fullmatch(r"\d{4}_(0[1-9]|1[0-2])", period):
        raise ValueError("Período inválido: usar AAAA_MM.")
    for case in cases:
        if not str(case.get("patient_name") or "").strip():
            raise ValueError("Cada expediente debe identificar al paciente.")
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("El informe debe tener extensión .xlsx.")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = f"Comparación documental {period}"
    wb.properties.creator = "CompareForms"
    wb.properties.subject = "Auditoría documental de expedientes y localización de diferencias"
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    context = f"Período {period} · Ejecución {run_id} · {generated}"
    has_incomplete = status != "completed" or any((c.get("comparison") or {}).get("status") not in {"with_differences", "no_differences_detected"} for c in cases)
    summary = _base(wb, "Resumen ejecutivo", "Comparación documental por paciente", context,
                    ["Paciente", "Resultado documental", "Observaciones", "Páginas añadidas", "Páginas eliminadas", "Páginas reubicadas", "Revisión requerida"], [38, 30, 16, 19, 19, 19, 31])
    summary.cell(5, 1, "PARCIAL / NO CONCLUYENTE" if has_incomplete else "COMPARACIÓN COMPLETADA").font = Font(name="Arial", size=10, bold=True, color="9C5700" if has_incomplete else NAVY)
    expedients = _base(wb, "Expedientes", "Identificación y alcance de la comparación", context,
                       ["Paciente", "Expediente", "Resultado", "Páginas original", "Páginas modificado", "Añadidas", "Eliminadas", "Reubicadas", "Observaciones", "Variación textual", "Prioridad", "Cobertura evaluada original", "Cobertura evaluada modificado", "Limitaciones"],
                       [38, 38, 30, 18, 18, 15, 15, 15, 17, 18, 18, 23, 23, 80])
    findings_ws = _base(wb, "Hallazgos", "Diferencias y contenido antes y después", context,
                        ["Paciente", "Tipo de cambio", "Categoría", "Página original", "Página modificado", "Contenido original", "Contenido modificado", "Observación documental", "Confianza de extracción", "Requiere revisión", "Evidencia", "Fragmento"],
                        [36, 18, 26, 18, 18, 65, 65, 70, 19, 20, 29, 13])
    findings_ws.freeze_panes = "C7"
    guide = _base(wb, "Guía de páginas", "Dónde localizar cada página", context,
                  ["Paciente", "Página original", "Página modificado", "Correspondencia", "Documento original", "Documento modificado", "Similitud de asociación", "Revisión requerida"],
                  [38, 19, 19, 24, 34, 34, 21, 21])
    money = _base(wb, "Importes", "Cambios de importes con contexto documental", context,
                  ["Paciente", "Concepto o línea original", "Concepto o línea modificada", "Página original", "Página modificado", "Posición en la línea", "Importe original (USD)", "Importe modificado (USD)", "Variación (USD)", "Variación relativa", "Evidencia", "Fragmento"],
                  [38, 65, 65, 18, 18, 17, 24, 24, 22, 22, 29, 13])
    review = _base(wb, "Revisión del auditor", "Registro de revisión de observaciones", "La edición de esta hoja no actualiza el portal. Registre la decisión definitiva en CompareForms.",
                   ["Paciente", "Evidencia", "Tipo detectado", "Página original", "Página modificado", "Decisión del auditor", "Observación del auditor"], [38, 29, 20, 19, 19, 28, 85])
    trace = _base(wb, "Trazabilidad", "Documentos y versión de análisis", context,
                  ["Paciente", "Expediente", "Documento original", "Documento modificado", "SHA-256 original", "SHA-256 modificado", "Motor", "Ejecución"],
                  [38, 38, 38, 38, 70, 70, 31, 42])
    for case in cases:
        patient = case["patient_name"]
        comp = case.get("comparison") or {}
        case_status = comp.get("status", case.get("status", "inconclusive"))
        completed = case_status in {"with_differences", "no_differences_detected"}
        fs = comp.get("findings", [])
        counts_known = comp.get("page_counts_reconciled", completed) and not any(p.get("review_required") for p in comp.get("page_map", []))
        count = lambda key: comp.get(key, "No evaluado") if counts_known else "No determinado"
        _append(summary, [patient, STATUS_LABELS.get(case_status, "No evaluado"), len(fs) if comp else "No evaluado", count("pages_added"), count("pages_removed"), count("pages_relocated"), "Sí: evidencia pendiente de revisión" if not completed else ("Sí: validar observaciones" if fs else "Sin pendientes del motor")])
        limit_parts = _parts("\n".join(comp.get("limitations", [])))
        for part_idx, limitation in enumerate(limit_parts):
            r = _append(expedients, [patient, case.get("id", ""), STATUS_LABELS.get(case_status, "No evaluado"), comp.get("pages_original", "No evaluado"), comp.get("pages_modified", "No evaluado"), count("pages_added"), count("pages_removed"), count("pages_relocated"), len(fs) if comp else "No evaluado", comp.get("text_change_ratio"), comp.get("priority") if comp.get("priority") is not None else "Sin calibrar", comp.get("pages_evaluated_original", "No evaluado"), comp.get("pages_evaluated_modified", "No evaluado"), (f"Continuación {part_idx + 1}/{len(limit_parts)}: " if len(limit_parts) > 1 else "") + limitation])
            expedients.cell(r, 10).number_format = "0.0%"
        for f in fs:
            kind = TYPE_LABELS.get(f["change_type"], "Revisión")
            before, after, descriptions = _parts(f.get("before")), _parts(f.get("after")), _parts(f.get("description"))
            chunks = max(len(before), len(after), len(descriptions))
            for k in range(chunks):
                r = _append(findings_ws, [patient, kind, f.get("category", "Contenido documental"), _page(f.get("page_original")), _page(f.get("page_modified")), before[k] if k < len(before) else "", after[k] if k < len(after) else "", descriptions[k] if k < len(descriptions) else "", f.get("confidence"), "Sí" if f.get("review_required") else "Validación de auditoría", f["id"], f"{k + 1}/{chunks}"])
                findings_ws.cell(r, 2).fill = PatternFill("solid", fgColor=COLORS[kind])
                findings_ws.cell(r, 9).number_format = "0%"
                _link(findings_ws, r, 4, case.get("original_id"), f.get("page_original"))
                _link(findings_ws, r, 5, case.get("modified_id"), f.get("page_modified"))
            _append(review, [patient, f["id"], kind, _page(f.get("page_original")), _page(f.get("page_modified")), "Pendiente", ""])
            for amount in f.get("monetary_changes", []):
                for k in range(max(len(before), len(after))):
                    # Amounts occur once; evidence continuations do not duplicate
                    # charge totals when the user filters or sums a column.
                    av = Decimal(amount["before"]) if k == 0 else None
                    bv = Decimal(amount["after"]) if k == 0 else None
                    r = _append(money, [patient, before[k] if k < len(before) else "", after[k] if k < len(after) else "", _page(f.get("page_original")), _page(f.get("page_modified")), amount["position"], av, bv, bv - av if k == 0 else None, (bv - av) / av if k == 0 and av else None, f["id"], f"{k + 1}/{max(len(before), len(after))}"])
                    for col in (7, 8, 9):
                        money.cell(r, col).number_format = '"$"#,##0.0000;[Red]-"$"#,##0.0000'
                    money.cell(r, 10).number_format = "0.00%"
        for pm in comp.get("page_map", []):
            _append(guide, [patient, _page(pm.get("page_original")), _page(pm.get("page_modified")), "Sin correspondencia (revisar)" if pm.get("review_required") and pm.get("status") in {"added", "removed"} else MAP_LABELS.get(pm.get("status"), "Por revisar"), pm.get("document_type_original", "Sin equivalente" if pm.get("page_original") is None else "Documento"), pm.get("document_type_modificado", "Sin equivalente" if pm.get("page_modified") is None else "Documento"), pm.get("similarity"), "Sí" if pm.get("review_required") else "No"])
            guide.cell(guide.max_row, 7).number_format = "0.0%"
            _link(guide, guide.max_row, 2, case.get("original_id"), pm.get("page_original"))
            _link(guide, guide.max_row, 3, case.get("modified_id"), pm.get("page_modified"))
        _append(trace, [patient, case.get("id", ""), case.get("original_id", ""), case.get("modified_id", ""), comp.get("hash_original", "No disponible"), comp.get("hash_modified", "No disponible"), comp.get("engine_version", "No disponible"), run_id])
    if review.max_row >= HEADER_ROW + 1:
        dv = DataValidation(type="list", formula1='"Pendiente,Confirmado,Falso positivo,Requiere revisión"', allow_blank=False)
        dv.errorTitle, dv.error, dv.showErrorMessage = "Decisión no válida", "Seleccione una decisión de la lista.", True
        review.add_data_validation(dv)
        dv.add(f"F7:F{review.max_row}")
        for row in review.iter_rows(min_row=7, min_col=6, max_col=7):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for sheet, name in ((summary, "ResumenPacientes"), (expedients, "ExpedientesComparados"), (findings_ws, "DiferenciasDocumentales"), (guide, "CorrespondenciaPaginas"), (money, "VariacionImportes"), (review, "RevisionAuditoria"), (trace, "TrazabilidadDocumental")):
        _finish(sheet, name)
    if not money.tables:
        money.cell(7, 1, "No se confirmaron cambios de importes con evidencia suficiente. Revise las limitaciones del expediente.")
        money.merge_cells("A7:F7")
        money.cell(7, 1).alignment = Alignment(wrap_text=True, vertical="top")
        money.row_dimensions[7].height = 36
    method = wb.create_sheet("Metodología")
    method.sheet_view.showGridLines = False
    method.column_dimensions["A"].width = 35
    method.column_dimensions["B"].width = 115
    method.cell(2, 1, "Cómo interpretar la comparación").font = Font(name="Arial", size=16, color=NAVY, bold=True)
    definitions = [
        ("Localización", "Los números son páginas físicas del contador del visor PDF, no números impresos dentro del formulario. Sin equivalente significa que esa página no tiene correspondencia automática en el otro PDF."),
        ("Tipos de cambio", "Añadido: aparece contenido o una página en el modificado. Eliminado: desaparece contenido o una página del original. Modificado: cambia contenido en páginas correspondientes. Reubicado: cambia el orden relativo. Revisión: no existe evidencia suficiente para una conclusión automática."),
        ("Filtrar observaciones", "En Hallazgos, use el filtro de Paciente y de Tipo de cambio. Las columnas Contenido original y Contenido modificado conservan la evidencia antes y después. Las continuaciones comparten el mismo ID de evidencia; no son hallazgos adicionales."),
        ("Correspondencia de páginas", "Opción A: asignación global uno a uno de costo mínimo con páginas sin pareja permitidas. Una incorporación que desplaza las páginas siguientes no se cuenta como reordenación. Las asociaciones cercanas o páginas ilegibles requieren revisión."),
        ("Conteo físico", "Páginas modificado menos páginas original debe ser igual a añadidas menos eliminadas. En resultados no concluyentes, las páginas sin correspondencia pueden requerir reclasificación como sustitución, división o fusión."),
        ("Variación textual", "Tokens normalizados (sin diferencias de mayúsculas o tildes) cambiados divididos por el máximo de tokens de cada pareja evaluada. Se suman ambos componentes sobre parejas legibles no ambiguas. No incluye páginas sin pareja. Los campos antes/después sí preservan mayúsculas y tildes. No es porcentaje de error clínico ni de gravedad."),
        ("Importes", "Solo se calculan importes explícitos con signo monetario, hasta 15 dígitos, contexto estable y texto nativo legible. Se conserva Decimal y precisión de fuente. No se confirman importes basados solo en OCR ni separadores ambiguos. No se suman facturas de soporte como cargos del paciente. Un origen cero no tiene variación relativa calculable."),
        ("Confianza", "Indica calidad de extracción o similitud, según la columna. La confianza de texto nativo es heurística, no probabilidad validada. No confundir confianza con prioridad o corrección médica."),
        ("Gráficos y firmas", "Se compara la representación de páginas a 144 dpi. Una diferencia gráfica puede ser maquetación, imagen, sello, firma o texto. No acredita autenticidad ni demuestra ausencia de una firma. Revisar visualmente los casos señalados."),
        ("Prioridad", "No se emiten puntajes 1–100 sin una rúbrica calibrada y aprobada por auditoría. Cero solo corresponde a comparación completa sin diferencias detectadas. No evaluado y sin calibrar no equivalen a cero."),
        ("Decisiones del auditor", "La hoja Revisión del auditor permite registrar decisiones como apoyo. Estas ediciones no se sincronizan con PostgreSQL; la decisión oficial y su historial se registran en el portal."),
        ("Fuente y causalidad", "El documento modificado es referencia aportada por auditoría, no verdad clínica infalible. Las diferencias no prueban su causa ni un defecto de Roboti. Un origen DALIA debe analizarse con su procedencia propia."),
        ("Alcance del análisis", "No se entrenó ni utilizó un modelo de aprendizaje. Se preservan los PDF originales. Las limitaciones de OCR, render o asociación se muestran en Expedientes. No concluyente exige revisión antes de afirmar que no hay diferencias."),
        ("Abrir los PDF", "Las páginas tienen enlaces al portal únicamente cuando COMPAREFORMS_PUBLIC_URL está configurado. El navegador requiere sesión autorizada. Sin enlaces, use Guía de páginas y el visor del expediente."),
    ]
    for r, (label, text) in enumerate(definitions, 5):
        method.cell(r, 1, label).font = Font(name="Arial", size=10, color=NAVY, bold=True)
        method.cell(r, 2, text).font = Font(name="Arial", size=10)
        method.cell(r, 1).alignment = method.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
        method.row_dimensions[r].height = max(42, math.ceil(len(text) / 110) * 14 + 12)
    # Place methodology before traceability, matching reader workflow.
    wb.move_sheet(method, offset=-1)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".writing.xlsx")
    if path.exists():
        raise FileExistsError("El informe ya existe; use una nueva versión de ejecución.")
    wb.save(temporary)
    # Validate the actual serialized file before publishing it.
    check = load_workbook(temporary, read_only=False, data_only=False)
    expected = {"Resumen ejecutivo", "Expedientes", "Hallazgos", "Guía de páginas", "Importes", "Revisión del auditor", "Metodología", "Trazabilidad"}
    if set(check.sheetnames) != expected or check["Hallazgos"].cell(6, 1).value != "Paciente":
        check.close()
        raise RuntimeError("Falló la validación estructural del informe.")
    check.close()
    temporary.replace(path)
    return path
