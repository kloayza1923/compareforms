from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.reporting import write_report


def case():
    return {"id": "case-1", "patient_name": "PACIENTE SINTETICO DE PRUEBA", "original_id": "951a2d69-6d7b-4979-9be8-d62cb8dbf80d", "modified_id": "d339af38-8bb0-49b5-87a7-c08bbddc1d9a", "status": "completed", "comparison": {
        "status": "with_differences", "engine_version": "test", "pages_original": 2, "pages_modified": 2, "pages_added": 1, "pages_removed": 1, "pages_relocated": 0,
        "page_counts_reconciled": True, "limitations": [], "findings": [
            {"id": "F1", "change_type": "removed", "category": "Página documental", "description": "Página original eliminada", "before": "Contenido anterior", "after": "", "page_original": 1, "page_modified": None, "confidence": 0.99, "review_required": False},
            {"id": "F2", "change_type": "modified", "category": "Importe documental", "description": "Importe de pinza modificado", "before": "Pinza $120,18", "after": "Pinza $118,75", "page_original": 2, "page_modified": 1, "confidence": 0.99, "review_required": False,
             "monetary_changes": [{"position": 1, "before": "120.18", "after": "118.75", "delta": "-1.43", "relative_change": "-0.0118"}]},
        ], "page_map": [{"page_original": 1, "page_modified": None, "status": "removed", "similarity": 0.1}, {"page_original": 2, "page_modified": 1, "status": "matched", "similarity": 0.98}, {"page_original": None, "page_modified": 2, "status": "added", "similarity": 0.1}]}}


def test_auditor_columns_and_filters(tmp_path):
    path = write_report(tmp_path / "report.xlsx", period="2026_06", run_id="synthetic-run", cases=[case()], status="completed")
    wb = load_workbook(path)
    assert len(wb.sheetnames) == 8
    ws = wb["Hallazgos"]
    assert [ws.cell(6, n).value for n in range(1, 9)] == ["Paciente", "Tipo de cambio", "Categoría", "Página original", "Página modificado", "Contenido original", "Contenido modificado", "Observación documental"]
    assert ws["A7"].value == "PACIENTE SINTETICO DE PRUEBA"
    assert ws["B7"].value == "Eliminado"
    assert ws["D7"].value == 1 and ws["E7"].value == "Sin equivalente"
    assert ws.freeze_panes == "C7"
    assert len(ws.tables) == 1
    assert ws.tables["DiferenciasDocumentales"].autoFilter.ref == "A6:L8"
    assert wb["Importes"]["G7"].value == 120.18
    assert wb["Importes"]["H7"].value == 118.75
    assert wb["Importes"]["I7"].value == -1.43
    assert len(wb["Revisión del auditor"].data_validations.dataValidation) == 1
    assert wb["Expedientes"]["K7"].value == "Sin calibrar"
    wb.close()


def test_empty_execution_never_generates_normal_excel(tmp_path):
    with pytest.raises(ValueError, match="NO_COMPARABLE_PAIRS"):
        write_report(tmp_path / "none.xlsx", period="2026_06", run_id="empty", cases=[], status="completed")
    assert not (tmp_path / "none.xlsx").exists()


def test_inconclusive_does_not_become_green_zero(tmp_path):
    c = case()
    c["comparison"].update(status="inconclusive", priority=None, page_counts_reconciled=False, limitations=["OCR no disponible"])
    path = write_report(tmp_path / "report.xlsx", period="2026_06", run_id="partial", cases=[c], status="completed")
    wb = load_workbook(path)
    assert wb["Resumen ejecutivo"]["A5"].value == "PARCIAL / NO CONCLUYENTE"
    assert wb["Resumen ejecutivo"]["B7"].value == "No concluyente"
    assert wb["Resumen ejecutivo"]["D7"].value == "No determinado"
    wb.close()


def test_formula_injection_is_literal(tmp_path):
    c = case()
    c["patient_name"] = '=HYPERLINK("https://attacker.invalid")'
    c["comparison"]["findings"][0]["before"] = "  =1+1"
    path = write_report(tmp_path / "safe.xlsx", period="2026_06", run_id="safe", cases=[c], status="completed")
    wb = load_workbook(path)
    assert wb["Hallazgos"]["A7"].data_type == "s"
    assert wb["Hallazgos"]["F7"].data_type == "s"
    assert wb["Hallazgos"]["F7"].value == "'  =1+1"
    wb.close()


def test_long_evidence_is_not_silently_truncated(tmp_path):
    c = case()
    original = "Evidencia documental completa con texto extenso. " * 900
    c["comparison"]["findings"][0]["before"] = original
    path = write_report(tmp_path / "long.xlsx", period="2026_06", run_id="long", cases=[c], status="completed")
    wb = load_workbook(path)
    ws = wb["Hallazgos"]
    reconstructed = "".join(ws.cell(r, 6).value or "" for r in range(7, ws.max_row + 1) if ws.cell(r, 11).value == "F1")
    assert reconstructed == original
    # Summary counts unique evidence, not continuation rows.
    assert wb["Resumen ejecutivo"]["C7"].value == 2
    wb.close()


def test_authenticated_links_when_public_url_is_configured(tmp_path, monkeypatch):
    monkeypatch.setenv("COMPAREFORMS_PUBLIC_URL", "https://compareforms.example.test")
    path = write_report(tmp_path / "links.xlsx", period="2026_06", run_id="links", cases=[case()], status="completed")
    wb = load_workbook(path)
    assert wb["Hallazgos"]["D7"].hyperlink.target.endswith("/api/v1/documents/951a2d69-6d7b-4979-9be8-d62cb8dbf80d/content#page=1")
    assert wb["Hallazgos"]["E7"].hyperlink is None
    wb.close()


def test_no_overwrite_of_previous_run(tmp_path):
    path = tmp_path / "report.xlsx"
    write_report(path, period="2026_06", run_id="one", cases=[case()], status="completed")
    before = path.read_bytes()
    with pytest.raises(FileExistsError):
        write_report(path, period="2026_06", run_id="two", cases=[case()], status="completed")
    assert path.read_bytes() == before


def test_unmatched_candidate_is_never_confirmed_removal_in_guide(tmp_path):
    c = case()
    c["comparison"]["page_map"][0]["review_required"] = True
    path = write_report(tmp_path / "review.xlsx", period="2026_06", run_id="review", cases=[c], status="partial")
    wb = load_workbook(path)
    assert wb["Guía de páginas"]["D7"].value == "Sin correspondencia (revisar)"
    assert wb["Resumen ejecutivo"]["D7"].value == "No determinado"
    wb.close()
