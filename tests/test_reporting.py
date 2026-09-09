from pathlib import Path

import csv

from openpyxl import load_workbook

from app.reporting import write_reports


def test_reports_are_created_and_xlsx_reopens(tmp_path: Path):
    comparison = {
        "index": 1, "name_before": "PERSONA", "identical": False, "name_matches": True,
        "before": {"filename": "1 - PERSONA.pdf", "sha256": "a" * 64, "size_bytes": 100, "page_count": 1, "pages_requiring_ocr": [], "pages_ocr_applied": [1]},
        "after": {"filename": "1 - PERSONA.pdf", "sha256": "b" * 64, "size_bytes": 110, "page_count": 2, "pages_requiring_ocr": [2], "pages_ocr_applied": []},
        "page_delta": 1, "size_delta_bytes": 10,
        "text_diff": {"tokens_added": 3, "tokens_deleted": 1, "change_percent": 20},
        "scores": {"technical_change_percent": 30, "severity": 80, "level": "muy alto"},
        "sensitive_changes": {key: {"added_count": 1 if key == "dates" else 0, "removed_count": 0} for key in ("amounts", "dates", "identifiers", "diagnostic_codes", "signature_terms")},
        "metadata_changes": {"producer": {"before": "a", "after": "b"}},
        "monetary_line_changes": [{"evidence_id": "E_MONEY_LINE_1", "status": "different_amounts", "page_before": 1, "page_after": 1, "description": "Pinza de Biopsia Estandar", "amounts_before": ["$95,00", "$104,50"], "amounts_after": ["$95,00", "$95,00"], "review_note": "Diferencia de importe en línea documental; requiere revisión contra los documentos fuente."}],
        "ollama": {"status": "ok", "model": "gemma4:e2b-it-qat", "usage": {"prompt_tokens": 10, "completion_tokens": 5, "total_tokens": 15}, "summary": {"summary": {"text": "Resumen literal.", "evidence_ids": ["E_PAGE_DELTA"]}, "main_changes": [], "alerts": [], "limitations": []}},
    }
    manifest = {"run_id": "run-test", "month": "2025_10", "generated_at": "2025-10-31T00:00:00+00:00", "scope": "selected_indices", "requested_indices": [1], "pair_count": 1, "failure_count": 0, "limitations": ["Revisión humana requerida."], "comparisons": [comparison]}
    paths = write_reports("2025_10", tmp_path, manifest)
    assert all(Path(value).is_file() for value in paths.values())
    workbook = load_workbook(paths["xlsx"], read_only=False)
    assert workbook.sheetnames == ["Resumen ejecutivo", "Comparativo", "Hallazgos", "Metodología", "Ejecución"]
    assert workbook["Resumen ejecutivo"]["A6"].value == "2025_10"
    assert workbook["Comparativo"]["A5"].value == 1
    assert workbook["Comparativo"].freeze_panes == "C5"
    assert workbook["Comparativo"]["F5"].number_format == "0.0%"
    assert workbook["Comparativo"].auto_filter.ref == "A4:L5"
    assert workbook["Hallazgos"].max_row >= 5
    assert workbook["Hallazgos"].freeze_panes == "A5"
    money_row = next(row for row in workbook["Hallazgos"].iter_rows(min_row=5, values_only=True) if row[2] == "Diferencia de importe en línea documental")
    assert money_row[3:9] == (1, 1, "Pinza de Biopsia Estandar", "$95,00, $104,50", "$95,00, $95,00", "E_MONEY_LINE_1")
    assert workbook["Ejecución"]["B6"].value == "run-test"
    assert workbook["Ejecución"]["C24"].value == "a" * 64
    workbook.close()
    with Path(paths["csv"]).open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.DictReader(stream))
    assert len(csv_rows) == 1
    assert csv_rows[0]["technical_change_percent"] == "30"
