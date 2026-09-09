from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfWriter

from app.config import Settings
from app.service import load_results, run_month


def make_pdf(path: Path, pages: int):
    writer = PdfWriter()
    for _ in range(pages): writer.add_blank_page(width=612, height=792)
    with path.open("wb") as stream: writer.write(stream)


def test_end_to_end_without_ollama_or_ocr(tmp_path: Path):
    month = tmp_path / "2025_10"; before = month / "pdf_origen"; after = month / "pdf_modificado"
    before.mkdir(parents=True); after.mkdir()
    make_pdf(before / "1 - PERSONA.pdf", 1); make_pdf(after / "1 - PERSONA.pdf", 2)
    settings = Settings(data_root=tmp_path, ocr_enabled=False)
    result = run_month("2025_10", settings, use_ollama=False, workers=1)
    assert result["pairs"] == 1 and result["failures"] == 0 and result["run_id"]
    assert load_results("2025_10", settings)["run_id"] == result["run_id"]
    workbook = load_workbook(result["artifacts"]["xlsx"], read_only=True)
    assert workbook.sheetnames == ["Resumen ejecutivo", "Comparativo", "Hallazgos", "Metodología", "Ejecución"]
    assert workbook["Comparativo"].max_row == 5
    workbook.close()
